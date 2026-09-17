"""Extrae las hojas 4.1 y 4.2 sin modificar ni recalcular el libro de Excel.

Requiere openpyxl. Ejecucion: python scripts/actualizar_escenarios_excel.py
    --workbook /ruta/Figuras.xlsx
Los valores del MFMP corresponden a los resultados guardados del MFMP 2024.
"""

import argparse
import csv
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    args = parser.parse_args()
    source = args.workbook.resolve()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    debt, balance = workbook["4.1"], workbook["4.2"]
    years = [debt.cell(2, column).value for column in range(2, 19)]
    assert years == list(range(2019, 2036)), "Revisar el horizonte de la hoja 4.1"
    assert years == [balance.cell(2, c).value for c in range(2, 19)]

    rows = []
    for column, year in enumerate(years, start=2):
        values = [sheet.cell(row, column).value
                  for sheet in (debt, balance) for row in (3, 4, 5)]
        assert all(isinstance(v, (int, float)) and math.isfinite(v)
                   for v in values), f"Valor ausente o no numerico en {year}"
        rows.append([year, *values])
        if year <= 2027:
            assert math.isclose(values[1], values[2], abs_tol=1e-9)
            assert math.isclose(values[4], values[5], abs_tol=1e-9)

        if year >= 2027:
            r, g = debt.cell(7, column).value, debt.cell(8, column).value
            factor = (1 + r) / (1 + g)
            for debt_row, balance_row in ((4, 4), (5, 5)):
                independent = (factor * debt.cell(debt_row, column - 1).value
                               - balance.cell(balance_row, column).value)
                assert math.isclose(debt.cell(debt_row, column).value,
                                    independent, abs_tol=1e-9), year
        if year >= 2028:
            independent = 0.2 + 0.1 * (debt.cell(5, column - 1).value - 55)
            assert math.isclose(balance.cell(5, column).value,
                                independent, abs_tol=1e-9), year

    destination = Path(__file__).resolve().parents[1] / "datos"
    destination.mkdir(exist_ok=True)
    table_path = destination / "escenarios_4_1_4_2.csv"
    with table_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["year", "deuda_mfmp2024", "deuda_actualizacion",
                         "deuda_consolidacion", "balance_mfmp2024",
                         "balance_actualizacion", "balance_consolidacion"])
        writer.writerows([[row[0], *[format(v, ".12g") for v in row[1:]]]
                          for row in rows])

    # Verificacion de los valores serializados, no solo de la matriz en memoria.
    with table_path.open(encoding="utf-8", newline="") as stream:
        serialized = list(csv.reader(stream))[1:]
    for saved, original in zip(serialized, rows, strict=True):
        assert int(saved[0]) == original[0]
        assert all(math.isclose(float(a), b, abs_tol=1e-8, rel_tol=1e-11)
                   for a, b in zip(saved[1:], original[1:], strict=True))

    metadata = {
        "source": source.name,
        "source_sha256": source_hash,
        "source_modified_utc": datetime.fromtimestamp(
            source.stat().st_mtime, timezone.utc).isoformat(),
        "units": "porcentaje del PIB",
        "ranges": {"deuda": "4.1!B2:R5", "balance_primario": "4.2!B2:R5"},
        "mfmp_vintage": 2024,
        "historical_through": 2025,
        "forecast_from": 2026,
        "source_values": "Valores guardados en Excel; no se recalculan vinculos externos",
        "forecast_assumptions": {
            "r": debt["J7"].value, "g": debt["J8"].value,
            "consolidation_from": 2028,
            "primary_balance_rule": "0.2 + 0.1 * (deuda del ano anterior - 55)"
        },
        "excluded": "Celdas S:T de 4.2 sin encabezados de ano"
    }
    (destination / "escenarios_4_1_4_2_fuente.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workbook.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    print(f"Verificados {len(rows)} anos y {len(rows) * 6} valores; Excel sin cambios.")
    print("2035:", json.dumps(dict(zip(
        ("year", "deuda_mfmp2024", "deuda_actualizacion", "deuda_consolidacion",
         "balance_mfmp2024", "balance_actualizacion", "balance_consolidacion"),
        rows[-1])), ensure_ascii=True))


if __name__ == "__main__":
    main()
